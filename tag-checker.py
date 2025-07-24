import os
import re
import tkinter as tk
from tkinter import ttk, filedialog, scrolledtext, messagebox
import threading
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
import json
import time
import pickle
import datetime

class ExcelSearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel File Search Tool")
        self.root.geometry("800x600")
        
        # Search state variables
        self.searching = False
        self.paused = False
        self.search_thread = None
        self.current_search_state = {
            'directories': [],
            'pattern': '',
            'include_subdirs': True,
            'progress': [],
            'results': [],
            'pending_files': [],
            'files_checked': 0,
            'total_matches': 0,
            'current_directory_index': 0,
            'last_file_path': ''
        }
        
        # Create main frames
        self.top_frame = ttk.Frame(root, padding="10")
        self.top_frame.pack(fill=tk.X)
        
        self.dir_frame = ttk.LabelFrame(root, padding="10", text="Search Directories")
        self.dir_frame.pack(fill=tk.X, padx=10)
        
        self.middle_frame = ttk.Frame(root, padding="10")
        self.middle_frame.pack(fill=tk.X)
        
        self.button_frame = ttk.Frame(root, padding="5")
        self.button_frame.pack(fill=tk.X)
        
        self.results_frame = ttk.Frame(root, padding="10")
        self.results_frame.pack(fill=tk.BOTH, expand=True)
        
        # Directory List
        self.dir_list_frame = ttk.Frame(self.dir_frame)
        self.dir_list_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.directories = []
        self.dir_listbox = tk.Listbox(self.dir_list_frame, width=60, height=5)
        self.dir_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        # Scrollbar for directory list
        dir_scrollbar = ttk.Scrollbar(self.dir_list_frame, orient=tk.VERTICAL, command=self.dir_listbox.yview)
        dir_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.dir_listbox.configure(yscrollcommand=dir_scrollbar.set)
        
        # Directory buttons
        self.dir_btn_frame = ttk.Frame(self.dir_frame)
        self.dir_btn_frame.pack(side=tk.RIGHT, padx=(10, 0))
        
        self.add_dir_btn = ttk.Button(self.dir_btn_frame, text="Add Directory", command=self.add_directory)
        self.add_dir_btn.pack(pady=2)
        
        self.remove_dir_btn = ttk.Button(self.dir_btn_frame, text="Remove Selected", command=self.remove_directory)
        self.remove_dir_btn.pack(pady=2)
        
        # Search pattern input
        ttk.Label(self.middle_frame, text="Search Pattern (regex):").pack(side=tk.LEFT, padx=(0, 5))
        self.pattern_var = tk.StringVar()
        self.pattern_entry = ttk.Entry(self.middle_frame, textvariable=self.pattern_var, width=40)
        self.pattern_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))
        
        # Add checkbox for including subdirectories
        self.include_subdirs_var = tk.BooleanVar(value=True)
        self.include_subdirs_check = ttk.Checkbutton(self.middle_frame, text="Include Subdirectories", variable=self.include_subdirs_var)
        self.include_subdirs_check.pack(side=tk.LEFT, padx=(5, 0))
        
        # Search control buttons
        self.search_btn = ttk.Button(self.button_frame, text="Search", command=self.start_search)
        self.search_btn.pack(side=tk.LEFT, padx=5)
        
        self.pause_resume_btn = ttk.Button(self.button_frame, text="Pause", command=self.toggle_pause_resume, state=tk.DISABLED)
        self.pause_resume_btn.pack(side=tk.LEFT, padx=5)
        
        self.stop_btn = ttk.Button(self.button_frame, text="Stop", command=self.stop_search, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=5)
        
        # Export button
        self.export_btn = ttk.Button(self.button_frame, text="Export Results", command=self.export_results, state=tk.NORMAL)
        self.export_btn.pack(side=tk.RIGHT, padx=5)
        
        # Add Clear and Load Project buttons
        self.clear_btn = ttk.Button(self.button_frame, text="Clear Project", command=self.clear_project)
        self.clear_btn.pack(side=tk.RIGHT, padx=5)
        
        self.load_project_btn = ttk.Button(self.button_frame, text="Load Project", command=self.load_project)
        self.load_project_btn.pack(side=tk.RIGHT, padx=5)
        
        self.save_project_btn = ttk.Button(self.button_frame, text="Save Project", command=self.save_project)
        self.save_project_btn.pack(side=tk.RIGHT, padx=5)
        
        # Progress and results
        ttk.Label(self.results_frame, text="Progress:").pack(anchor=tk.W)
        self.progress_text = scrolledtext.ScrolledText(self.results_frame, height=5, width=80)
        self.progress_text.pack(fill=tk.X, pady=(0, 10))
        
        ttk.Label(self.results_frame, text="Results:").pack(anchor=tk.W)
        self.results_text = scrolledtext.ScrolledText(self.results_frame, height=20, width=80)
        self.results_text.pack(fill=tk.BOTH, expand=True)
        
        # Status bar
        self.status_var = tk.StringVar()
        self.status_var.set("Ready")
        self.status_bar = ttk.Label(root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Set up window close handler
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
        
        # Load previous session if available
        self.session_file = "excel_search_session.pickle"
        self.load_session()
    
    def add_directory(self):
        directory = filedialog.askdirectory()
        if directory and directory not in self.directories:
            self.directories.append(directory)
            self.dir_listbox.insert(tk.END, directory)
    
    def remove_directory(self):
        selected_idx = self.dir_listbox.curselection()
        if selected_idx:
            idx = selected_idx[0]
            self.dir_listbox.delete(idx)
            self.directories.pop(idx)
    
    def toggle_pause_resume(self):
        if self.paused:
            # Resume
            self.paused = False
            self.pause_resume_btn.config(text="Pause")
            self.status_var.set("Resuming search...")
            self.log_progress("Resuming search...")
            
            # Start a new thread for resumed search
            self.search_thread = threading.Thread(target=self.resume_search)
            self.search_thread.daemon = True
            self.search_thread.start()
        else:
            # Pause
            self.paused = True
            self.pause_resume_btn.config(text="Resume")
            self.status_var.set("Search paused")
            self.log_progress("Search paused.")
    
    def stop_search(self):
        if self.searching:
            self.searching = False
            self.paused = False
            self.pause_resume_btn.config(text="Pause", state=tk.DISABLED)
            self.stop_btn.config(state=tk.DISABLED)
            self.search_btn.config(state=tk.NORMAL)
            self.status_var.set("Search stopped")
            self.log_progress("Search stopped.")
            
            # Clear current search state
            self.current_search_state = {
                'directories': self.directories.copy(),
                'pattern': self.pattern_var.get().strip(),
                'include_subdirs': self.include_subdirs_var.get(),
                'progress': [],
                'results': [],
                'pending_files': [],
                'files_checked': 0,
                'total_matches': 0,
                'current_directory_index': 0,
                'last_file_path': ''
            }
    
    def export_results(self):
        """Export search results to an Excel file"""
        if not self.current_search_state['results']:
            messagebox.showinfo("Export", "No results to export!")
            return
        
        # Ask user for save location
        timestamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        default_filename = f"search_results_{timestamp}.xlsx"
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=default_filename
        )
        
        if not file_path:
            return  # User cancelled
        
        try:
            self.status_var.set("Exporting results...")
            
            # Create a new workbook and add worksheets
            workbook = openpyxl.Workbook()
            
            # Remove default sheet and create our own
            default_sheet = workbook.active
            workbook.remove(default_sheet)
            
            # Create summary sheet
            summary_sheet = workbook.create_sheet(title="Summary")
            
            # Create results sheet
            results_sheet = workbook.create_sheet(title="Search Results")
            
            # Style definitions
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin'), 
                right=Side(style='thin'), 
                top=Side(style='thin'), 
                bottom=Side(style='thin')
            )
            
            # Summary sheet
            summary_sheet.append(["Excel Search Results"])
            summary_sheet.append([])
            summary_sheet.append(["Search Pattern:", self.current_search_state['pattern']])
            
            directories_str = ", ".join(self.current_search_state['directories'])
            summary_sheet.append(["Search Directories:", directories_str])
            
            summary_sheet.append(["Include Subdirectories:", "Yes" if self.current_search_state['include_subdirs'] else "No"])
            summary_sheet.append(["Files Checked:", self.current_search_state['files_checked']])
            summary_sheet.append(["Total Matches:", self.current_search_state['total_matches']])
            summary_sheet.append(["Date/Time:", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
            
            # Format summary sheet
            summary_sheet.column_dimensions['A'].width = 20
            summary_sheet.column_dimensions['B'].width = 60
            
            for row in range(1, 9):
                for col in range(1, 3):
                    cell = summary_sheet.cell(row=row, column=col)
                    if row == 1:
                        cell.font = Font(bold=True, size=14)
                    elif row >= 3:
                        if col == 1:
                            cell.font = Font(bold=True)
            
            # Results sheet - Headers
            results_sheet.append(["File Path", "Sheet Name", "Cell", "Value"])
            
            # Format headers
            for col in range(1, 5):
                cell = results_sheet.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
            
            # Set column widths
            results_sheet.column_dimensions['A'].width = 50  # File path
            results_sheet.column_dimensions['B'].width = 20  # Sheet name
            results_sheet.column_dimensions['C'].width = 10  # Cell
            results_sheet.column_dimensions['D'].width = 50  # Value
            
            # Add data rows
            row = 2
            for result in self.current_search_state['results']:
                file_path = result['file_path']
                
                for match in result['matches']:
                    sheet_name = match['sheet']
                    cell_addr = match['cell']
                    cell_value = match['value']
                    
                    results_sheet.append([file_path, sheet_name, cell_addr, cell_value])
                    
                    # Apply borders
                    for col in range(1, 5):
                        results_sheet.cell(row=row, column=col).border = thin_border
                    
                    row += 1
            
            # Save the workbook
            workbook.save(file_path)
            self.status_var.set(f"Results exported to {file_path}")
            self.log_progress(f"Results successfully exported to {file_path}")
            
            # Ask if the user wants to open the file
            if messagebox.askyesno("Export Complete", f"Results exported to {file_path}.\nDo you want to open the file now?"):
                os.startfile(file_path)
                
        except Exception as e:
            self.log_progress(f"Export error: {str(e)}")
            self.status_var.set("Export failed. See progress log for details.")
            messagebox.showerror("Export Error", f"Failed to export results: {str(e)}")
    
    def start_search(self):
        if self.searching:
            return
        
        if not self.directories:
            self.status_var.set("Error: Please add at least one search directory")
            return
        
        search_pattern = self.pattern_var.get().strip()
        
        if not search_pattern:
            self.status_var.set("Error: Please provide a search pattern")
            return
        
        try:
            re.compile(search_pattern)
        except re.error:
            self.status_var.set("Error: Invalid regular expression")
            return
        
        # Clear previous results
        self.progress_text.delete(1.0, tk.END)
        self.results_text.delete(1.0, tk.END)
        
        # Set up search state
        include_subdirs = self.include_subdirs_var.get()
        self.current_search_state = {
            'directories': self.directories.copy(),
            'pattern': search_pattern,
            'include_subdirs': include_subdirs,
            'progress': [],
            'results': [],
            'pending_files': [],
            'files_checked': 0,
            'total_matches': 0,
            'current_directory_index': 0,
            'last_file_path': ''
        }
        
        # Start search in a separate thread
        self.searching = True
        self.paused = False
        self.search_btn.config(state=tk.DISABLED)
        self.pause_resume_btn.config(state=tk.NORMAL, text="Pause")
        self.stop_btn.config(state=tk.NORMAL)
        self.status_var.set("Searching...")
        
        self.search_thread = threading.Thread(target=self.search_excel_files)
        self.search_thread.daemon = True
        self.search_thread.start()
    
    def resume_search(self):
        # Continue from where we left off
        try:
            pattern = re.compile(self.current_search_state['pattern'])
            
            # Resume processing any pending files
            for file_info in self.current_search_state['pending_files']:
                if not self.searching or self.paused:
                    return
                
                directory, file = file_info
                self._process_file(directory, file, pattern)
            
            # Clear pending files after processing
            self.current_search_state['pending_files'] = []
            
            # Continue with remaining directories
            self.search_excel_files(resume=True)
                
        except Exception as e:
            self.log_progress(f"Resume error: {str(e)}")
            self.root.after(0, lambda: self.status_var.set("Resume failed. See progress log for details."))
            self.searching = False
            self.paused = False
            self.root.after(0, lambda: self.search_btn.config(state=tk.NORMAL))
            self.root.after(0, lambda: self.pause_resume_btn.config(state=tk.DISABLED))
            self.root.after(0, lambda: self.stop_btn.config(state=tk.DISABLED))
    
    def search_excel_files(self, resume=False):
        try:
            pattern = re.compile(self.current_search_state['pattern'])
            
            if not resume:
                paths_str = ", ".join(self.current_search_state['directories'])
                self.log_progress(f"Starting search in {paths_str} for pattern: {self.current_search_state['pattern']}")
            else:
                self.log_progress(f"Resuming search from directory index {self.current_search_state['current_directory_index']}")
            
            # Start from the current directory index if resuming
            start_idx = self.current_search_state['current_directory_index']
            for idx, search_path in enumerate(self.current_search_state['directories'][start_idx:], start_idx):
                # Update current directory index
                self.current_search_state['current_directory_index'] = idx
                
                if not self.searching:
                    break
                
                self.log_progress(f"Searching in directory: {search_path}")
                
                if self.current_search_state['include_subdirs']:
                    for root, _, files in os.walk(search_path):
                        # Store files for current directory to handle pause/resume
                        file_batch = [(root, f) for f in files if f.endswith(('.xlsx', '.xlsm')) and not f.startswith('~$')]
                        self.current_search_state['pending_files'] = file_batch
                        
                        for directory, file in file_batch:
                            if not self.searching:
                                break
                                
                            if self.paused:
                                # Save current state and return
                                self.save_session()
                                return
                            
                            self._process_file(directory, file, pattern)
                        
                        # Clear pending files after processing this batch
                        self.current_search_state['pending_files'] = []
                        
                        if not self.searching or self.paused:
                            break
                else:
                    files = [f for f in os.listdir(search_path) 
                             if os.path.isfile(os.path.join(search_path, f)) and
                             f.endswith(('.xlsx', '.xlsm')) and not f.startswith('~$')]
                    
                    # Store files for current directory to handle pause/resume
                    file_batch = [(search_path, f) for f in files]
                    self.current_search_state['pending_files'] = file_batch
                    
                    for directory, file in file_batch:
                        if not self.searching:
                            break
                            
                        if self.paused:
                            # Save current state and return
                            self.save_session()
                            return
                        
                        self._process_file(directory, file, pattern)
                    
                    # Clear pending files after processing this batch
                    self.current_search_state['pending_files'] = []
            
            if self.searching and not self.paused:
                # Search completed
                self.log_progress(f"Search complete. Checked {self.current_search_state['files_checked']} files. Found {self.current_search_state['total_matches']} total matches.")
                self.root.after(0, lambda: self.status_var.set(f"Search complete. Found {self.current_search_state['total_matches']} matches."))
                
                # Reset search state
                self.searching = False
                self.root.after(0, lambda: self.search_btn.config(state=tk.NORMAL))
                self.root.after(0, lambda: self.pause_resume_btn.config(state=tk.DISABLED))
                self.root.after(0, lambda: self.stop_btn.config(state=tk.DISABLED))
        except Exception as e:
            self.log_progress(f"Search error: {str(e)}")
            self.root.after(0, lambda: self.status_var.set("Search failed. See progress log for details."))
            self.searching = False
            self.paused = False
            self.root.after(0, lambda: self.search_btn.config(state=tk.NORMAL))
            self.root.after(0, lambda: self.pause_resume_btn.config(state=tk.DISABLED))
            self.root.after(0, lambda: self.stop_btn.config(state=tk.DISABLED))
    
    def _process_file(self, directory, file, pattern):
        file_path = os.path.join(directory, file)
        self.current_search_state['last_file_path'] = file_path
        self.log_progress(f"Checking file: {file_path}")
        
        try:
            matches = self.search_in_excel(file_path, pattern)
            self.current_search_state['files_checked'] += 1
            
            if matches:
                self.current_search_state['total_matches'] += len(matches)
                self.log_progress(f"Found {len(matches)} matches in {file_path}")
                
                # Add results to results text
                result_text = f"\n--- {file_path} ---\n"
                self.log_result(result_text)
                
                # Store result in current search state
                result_entry = {"file_path": file_path, "matches": []}
                
                for match in matches:
                    sheet_name, cell_addr, cell_value = match
                    match_text = f"Sheet: {sheet_name}, Cell: {cell_addr}, Value: {cell_value}\n"
                    self.log_result(match_text)
                    
                    # Store match in current search state
                    result_entry["matches"].append({
                        "sheet": sheet_name,
                        "cell": cell_addr,
                        "value": cell_value
                    })
                
                # Add completed result entry to results list
                self.current_search_state['results'].append(result_entry)
            else:
                self.log_progress(f"No matches found in {file_path}")
        except Exception as e:
            self.log_progress(f"Error processing {file_path}: {str(e)}")
    
    def search_in_excel(self, file_path, pattern):
        matches = []
        
        # Open workbook in read-only mode
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            
            # Iterate through all cells with values
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cell_value = str(cell.value)
                        if pattern.search(cell_value):
                            cell_addr = f"{cell.column_letter}{cell.row}"
                            matches.append((sheet_name, cell_addr, cell_value))
                            
                    # Check if search has been paused or stopped
                    if not self.searching or self.paused:
                        workbook.close()
                        return matches
        
        workbook.close()
        return matches
    
    def log_progress(self, message):
        # Store in progress history
        self.current_search_state['progress'].append(message)
        self.root.after(0, lambda: self._append_text(self.progress_text, message + "\n"))
    
    def log_result(self, message):
        self.root.after(0, lambda: self._append_text(self.results_text, message))
    
    def _append_text(self, text_widget, message):
        text_widget.configure(state=tk.NORMAL)
        text_widget.insert(tk.END, message)
        text_widget.see(tk.END)
        text_widget.configure(state=tk.NORMAL)
    
    def save_session(self, custom_file=None):
        """Save the current search state to a file"""
        try:
            save_path = custom_file or self.session_file
            
            # Create a complete session state
            session_state = {
                'search_state': self.current_search_state,
                'directories': self.directories,
                'pattern': self.pattern_var.get(),
                'include_subdirs': self.include_subdirs_var.get(),
                'searching': self.searching,
                'paused': self.paused
            }
            
            with open(save_path, 'wb') as f:
                pickle.dump(session_state, f)
            
            if custom_file:
                self.log_progress(f"Project saved to {custom_file}")
                self.status_var.set(f"Project saved to {custom_file}")
            elif self.paused:
                self.log_progress("Session saved. You can close the application and resume later.")
                
        except Exception as e:
            self.log_progress(f"Error saving session: {str(e)}")
            if custom_file:
                messagebox.showerror("Save Error", f"Failed to save project: {str(e)}")
    
    def load_session(self):
        """Load the previous search session if available"""
        try:
            if os.path.exists(self.session_file):
                with open(self.session_file, 'rb') as f:
                    session_state = pickle.load(f)
                
                # Restore application state
                self.current_search_state = session_state['search_state']
                self.directories = session_state['directories']
                self.pattern_var.set(session_state['pattern'])
                self.include_subdirs_var.set(session_state['include_subdirs'])
                
                # Update directories listbox
                for directory in self.directories:
                    self.dir_listbox.insert(tk.END, directory)
                
                # Restore progress and results
                for progress_msg in self.current_search_state['progress']:
                    self._append_text(self.progress_text, progress_msg + "\n")
                
                for result in self.current_search_state['results']:
                    self._append_text(self.results_text, f"\n--- {result['file_path']} ---\n")
                    for match in result['matches']:
                        self._append_text(self.results_text, 
                            f"Sheet: {match['sheet']}, Cell: {match['cell']}, Value: {match['value']}\n")
                
                # Ask if user wants to resume the paused search
                if session_state['paused'] and session_state['searching']:
                    if messagebox.askyesno("Resume Search", 
                                          "A search was paused in your previous session. Do you want to resume it?"):
                        self.searching = True
                        self.paused = False
                        
                        self.search_btn.config(state=tk.DISABLED)
                        self.pause_resume_btn.config(state=tk.NORMAL, text="Pause")
                        self.stop_btn.config(state=tk.NORMAL)
                        
                        self.status_var.set("Resuming previous search...")
                        self.log_progress("Resuming search from previous session...")
                        
                        self.search_thread = threading.Thread(target=self.resume_search)
                        self.search_thread.daemon = True
                        self.search_thread.start()
                    else:
                        # User declined to resume, clean up the session
                        self.current_search_state['pending_files'] = []
                        self.searching = False
                        self.paused = False
                
                self.log_progress("Previous session loaded.")
                self.status_var.set("Previous session loaded.")
        except Exception as e:
            self.log_progress(f"Error loading session: {str(e)}")
            # If loading failed, start fresh
            self.current_search_state = {
                'directories': [],
                'pattern': '',
                'include_subdirs': True,
                'progress': [],
                'results': [],
                'pending_files': [],
                'files_checked': 0,
                'total_matches': 0,
                'current_directory_index': 0,
                'last_file_path': ''
            }
    
    def on_closing(self):
        if self.searching:
            if messagebox.askyesno("Quit", "A search is in progress. Do you want to save and quit?"):
                # Force pause state to save correctly
                if not self.paused:
                    self.paused = True
                self.save_session()
                self.root.destroy()
        else:
            self.root.destroy()
    
    def clear_project(self):
        """Clear all search data and reset the application state"""
        if self.searching and not messagebox.askyesno("Clear Project", 
                                                   "A search is in progress. Are you sure you want to clear everything?"):
            return
            
        # Stop any ongoing search
        if self.searching:
            self.stop_search()
            
        # Clear directories
        self.directories = []
        self.dir_listbox.delete(0, tk.END)
        
        # Clear search pattern
        self.pattern_var.set("")
        
        # Clear text areas
        self.progress_text.configure(state=tk.NORMAL)
        self.progress_text.delete(1.0, tk.END)
        self.progress_text.configure(state=tk.NORMAL)
        
        self.results_text.configure(state=tk.NORMAL)
        self.results_text.delete(1.0, tk.END)
        self.results_text.configure(state=tk.NORMAL)
        
        # Reset search state
        self.current_search_state = {
            'directories': [],
            'pattern': '',
            'include_subdirs': True,
            'progress': [],
            'results': [],
            'pending_files': [],
            'files_checked': 0,
            'total_matches': 0,
            'current_directory_index': 0,
            'last_file_path': ''
        }
        
        self.status_var.set("Project cleared")
        self.log_progress("Project cleared. All data has been reset.")
    
    def load_project(self):
        """Load a saved project file"""
        if self.searching and not messagebox.askyesno("Load Project", 
                                                   "A search is in progress. Are you sure you want to load a project?"):
            return
            
        # Stop any ongoing search
        if self.searching:
            self.stop_search()
            
        # Ask for project file
        project_file = filedialog.askopenfilename(
            defaultextension=".exs",
            filetypes=[("Excel Search Project", "*.exs"), ("Pickle Files", "*.pickle"), ("All Files", "*.*")],
            title="Load Project File"
        )
        
        if not project_file:
            return  # User cancelled
            
        try:
            with open(project_file, 'rb') as f:
                session_state = pickle.load(f)
                
            # Clear current state first
            self.clear_project()
                
            # Restore application state
            self.current_search_state = session_state['search_state']
            self.directories = session_state['directories']
            self.pattern_var.set(session_state['pattern'])
            self.include_subdirs_var.set(session_state['include_subdirs'])
            
            # Update directories listbox
            for directory in self.directories:
                self.dir_listbox.insert(tk.END, directory)
            
            # Restore progress and results
            for progress_msg in self.current_search_state['progress']:
                self._append_text(self.progress_text, progress_msg + "\n")
            
            for result in self.current_search_state['results']:
                self._append_text(self.results_text, f"\n--- {result['file_path']} ---\n")
                for match in result['matches']:
                    self._append_text(self.results_text, 
                        f"Sheet: {match['sheet']}, Cell: {match['cell']}, Value: {match['value']}\n")
            
            self.status_var.set(f"Project loaded from {project_file}")
            self.log_progress(f"Project loaded from {project_file}")
            
            # Ask if user wants to resume any paused search
            if session_state.get('paused', False) and session_state.get('searching', False):
                if messagebox.askyesno("Resume Search", 
                                      "This project contains a paused search. Do you want to resume it?"):
                    self.searching = True
                    self.paused = False
                    
                    self.search_btn.config(state=tk.DISABLED)
                    self.pause_resume_btn.config(state=tk.NORMAL, text="Pause")
                    self.stop_btn.config(state=tk.NORMAL)
                    
                    self.status_var.set("Resuming search...")
                    self.log_progress("Resuming search from loaded project...")
                    
                    self.search_thread = threading.Thread(target=self.resume_search)
                    self.search_thread.daemon = True
                    self.search_thread.start()
                
        except Exception as e:
            self.log_progress(f"Error loading project: {str(e)}")
            messagebox.showerror("Load Error", f"Failed to load project: {str(e)}")
            self.status_var.set("Failed to load project")
    
    def save_project(self):
        """Save the current project to a custom file"""
        if not self.directories and not self.current_search_state['results']:
            messagebox.showinfo("Save Project", "There's no project data to save!")
            return
            
        # Ask for save location
        timestamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        default_filename = f"excel_search_project_{timestamp}.exs"
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".exs",
            filetypes=[("Excel Search Project", "*.exs"), ("Pickle Files", "*.pickle")],
            initialfile=default_filename,
            title="Save Project As"
        )
        
        if not file_path:
            return  # User cancelled
            
        # Force pause if searching
        was_paused = self.paused
        if self.searching and not self.paused:
            self.toggle_pause_resume()
            
        # Save to the selected file
        self.save_session(custom_file=file_path)
        
        # Resume if we paused it for saving
        if self.searching and not was_paused and self.paused:
            self.toggle_pause_resume()

if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelSearchApp(root)
    root.mainloop()

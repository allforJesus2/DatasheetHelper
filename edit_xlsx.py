import os
import openpyxl
import xlwings as xw
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import json
import sys
import io
import threading
import queue

class ConfirmationDialog:
    def __init__(self, parent, existing_value, proposed_value, cell_address, sheet_name):
        self.result = None
        self.apply_to_all = False
        
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Confirm Value Replacement")
        self.dialog.geometry("600x450")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()  # Make dialog modal
        
        # Center the dialog
        self.dialog.update_idletasks()
        x = (parent.winfo_screenwidth() // 2) - (600 // 2)
        y = (parent.winfo_screenheight() // 2) - (450 // 2)
        self.dialog.geometry(f"600x450+{x}+{y}")
        
        # Create the dialog content
        self.create_widgets(existing_value, proposed_value, cell_address, sheet_name)
        
        # Wait for user response
        parent.wait_window(self.dialog)
    
    def create_widgets(self, existing_value, proposed_value, cell_address, sheet_name):
        # Main frame
        main_frame = ttk.Frame(self.dialog, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Title
        title_label = ttk.Label(main_frame, text="Value Replacement Confirmation", 
                               font=("Arial", 12, "bold"))
        title_label.pack(pady=(0, 20))
        
        # Location info
        location_frame = ttk.Frame(main_frame)
        location_frame.pack(fill=tk.X, pady=(0, 15))
        ttk.Label(location_frame, text=f"Sheet: {sheet_name}").pack(anchor=tk.W)
        ttk.Label(location_frame, text=f"Cell: {cell_address}").pack(anchor=tk.W)
        
        # Values comparison
        values_frame = ttk.LabelFrame(main_frame, text="Values", padding="10")
        values_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))
        
        # Existing value
        existing_frame = ttk.Frame(values_frame)
        existing_frame.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(existing_frame, text="Existing Value:", font=("Arial", 10, "bold")).pack(anchor=tk.W)
        existing_text = tk.Text(existing_frame, height=3, wrap=tk.WORD, state=tk.DISABLED)
        existing_text.pack(fill=tk.X, pady=(5, 0))
        existing_text.config(state=tk.NORMAL)
        existing_text.insert(tk.END, str(existing_value) if existing_value is not None else "(empty)")
        existing_text.config(state=tk.DISABLED)
        
        # Proposed value
        proposed_frame = ttk.Frame(values_frame)
        proposed_frame.pack(fill=tk.X)
        ttk.Label(proposed_frame, text="Proposed Value:", font=("Arial", 10, "bold")).pack(anchor=tk.W)
        proposed_text = tk.Text(proposed_frame, height=3, wrap=tk.WORD, state=tk.DISABLED)
        proposed_text.pack(fill=tk.X, pady=(5, 0))
        proposed_text.config(state=tk.NORMAL)
        proposed_text.insert(tk.END, str(proposed_value))
        proposed_text.config(state=tk.DISABLED)
        
        # Apply to all checkbox
        self.apply_to_all_var = tk.BooleanVar()
        apply_check = ttk.Checkbutton(main_frame, text="Apply this decision to all future conflicts", 
                                     variable=self.apply_to_all_var)
        apply_check.pack(pady=(0, 15))
        
        # Buttons
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X)
        
        ttk.Button(button_frame, text="Replace", command=self.replace, 
                  style="Accent.TButton").pack(side=tk.RIGHT, padx=(5, 0))
        ttk.Button(button_frame, text="Skip", command=self.skip).pack(side=tk.RIGHT, padx=(5, 0))
        ttk.Button(button_frame, text="Cancel All", command=self.cancel_all).pack(side=tk.LEFT)
    
    def replace(self):
        self.result = "replace"
        self.apply_to_all = self.apply_to_all_var.get()
        self.dialog.destroy()
    
    def skip(self):
        self.result = "skip"
        self.apply_to_all = self.apply_to_all_var.get()
        self.dialog.destroy()
    
    def cancel_all(self):
        self.result = "cancel"
        self.dialog.destroy()

class ExcelEditorApp:
    def __init__(self, master):
        self.master = master
        self.master.title("Excel Editor")
        self.master.geometry("900x700")  # Made taller to accommodate new section
        self.output_queue = queue.Queue()  # Add this line
        self.output_popup = None  # Track the output popup window
        self.edits = []
        self.replacements = []  # Now stores tuples (search, replace, exact_match)
        self.confirmation_settings = {
            'replace_all': None,  # None = ask, True = always replace, False = always skip
            'skip_all': False
        }
        self.create_widgets()

    def should_replace_value(self, existing_value, proposed_value, cell_address, sheet_name):
        """
        Determine if a value should be replaced based on confirmation settings and user input.
        Returns True if the value should be replaced, False otherwise.
        """
        # If there's no existing value (None or empty), always replace
        if existing_value is None or str(existing_value).strip() == "":
            return True
        
        # If the values are the same, no need to replace
        if str(existing_value) == str(proposed_value):
            self.output_queue.put(f"SKIPPED: '{existing_value}' → '{proposed_value}' (values are identical) at {sheet_name}!{cell_address}")
            return False
        
        # Check if we have a global setting
        if self.confirmation_settings['replace_all'] is True:
            return True
        elif self.confirmation_settings['replace_all'] is False:
            return False
        
        # If we're supposed to skip all, return False
        if self.confirmation_settings['skip_all']:
            return False
        
        # Show confirmation dialog
        dialog = ConfirmationDialog(self.master, existing_value, proposed_value, cell_address, sheet_name)
        
        if dialog.result == "cancel":
            # Cancel all future operations
            self.confirmation_settings['skip_all'] = True
            return False
        elif dialog.result == "replace":
            if dialog.apply_to_all:
                self.confirmation_settings['replace_all'] = True
            return True
        elif dialog.result == "skip":
            if dialog.apply_to_all:
                self.confirmation_settings['replace_all'] = False
            return False
        
        # Default to skip if something unexpected happens
        return False

    def create_widgets(self):
        # Add menu bar
        menubar = tk.Menu(self.master)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Import Edits", command=self.import_edits)
        file_menu.add_command(label="Export Edits", command=self.export_edits)
        menubar.add_cascade(label="File", menu=file_menu)
        self.master.config(menu=menubar)

        # File selection frame
        file_frame = ttk.LabelFrame(self.master, text="File Selection")
        file_frame.pack(padx=10, pady=10, fill=tk.X)

        self.file_selection = tk.StringVar(value="WALK")
        ttk.Radiobutton(file_frame, text="Walk Directory", variable=self.file_selection, value="WALK").pack(
            side=tk.LEFT)
        ttk.Radiobutton(file_frame, text="Single Folder", variable=self.file_selection, value="FOLDER").pack(
            side=tk.LEFT)
        ttk.Radiobutton(file_frame, text="Single/Multiple Files", variable=self.file_selection, value="FILES").pack(side=tk.LEFT)

        # Search/Replace frame
        replace_frame = ttk.LabelFrame(self.master, text="Search and Replace")
        replace_frame.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)

        # Search/Replace table
        self.replace_tree = ttk.Treeview(replace_frame, columns=("Search", "Replace", "Exact"), show="headings", height=5)
        self.replace_tree.heading("Search", text="Search Text")
        self.replace_tree.heading("Replace", text="Replace Text")
        self.replace_tree.heading("Exact", text="Exact Match")
        self.replace_tree.column("Exact", width=80, anchor=tk.CENTER)
        self.replace_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.replace_tree.bind("<Double-1>", self.edit_replace_pair)

        replace_scrollbar = ttk.Scrollbar(replace_frame, orient=tk.VERTICAL, command=self.replace_tree.yview)
        replace_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.replace_tree.configure(yscrollcommand=replace_scrollbar.set)

        # Search/Replace input frame
        replace_input_frame = ttk.Frame(self.master)
        replace_input_frame.pack(padx=10, pady=5, fill=tk.X)

        ttk.Label(replace_input_frame, text="Search Text:").grid(row=0, column=0)
        self.search_entry = ttk.Entry(replace_input_frame)
        self.search_entry.grid(row=0, column=1)

        ttk.Label(replace_input_frame, text="Replace Text:").grid(row=0, column=2)
        self.replace_entry = ttk.Entry(replace_input_frame)
        self.replace_entry.grid(row=0, column=3)

        self.exact_match_var = tk.BooleanVar()
        exact_check = ttk.Checkbutton(replace_input_frame, text="Exact Match", variable=self.exact_match_var)
        exact_check.grid(row=0, column=6, padx=5)

        ttk.Button(replace_input_frame, text="Add Replace Pair", command=self.add_replace_pair).grid(row=0, column=4, padx=5)
        ttk.Button(replace_input_frame, text="Delete Selected", command=self.delete_replace_pair).grid(row=0, column=5, padx=5)

        # Existing edits frame
        edits_frame = ttk.LabelFrame(self.master, text="Offset Edits")
        edits_frame.pack(padx=10, pady=10, fill=tk.BOTH, expand=True)

        # Edits table
        self.edits_tree = ttk.Treeview(edits_frame, columns=("Keyword", "Col Offset", "Row Offset", "New Value"),
                                       show="headings", height=5)
        self.edits_tree.heading("Keyword", text="Keyword")
        self.edits_tree.heading("Col Offset", text="Col Offset")
        self.edits_tree.heading("Row Offset", text="Row Offset")
        self.edits_tree.heading("New Value", text="New Value")
        self.edits_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.edits_tree.bind("<Double-1>", self.edit_offset_edit)

        scrollbar = ttk.Scrollbar(edits_frame, orient=tk.VERTICAL, command=self.edits_tree.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.edits_tree.configure(yscrollcommand=scrollbar.set)

        # Edit input frame
        input_frame = ttk.Frame(self.master)
        input_frame.pack(padx=10, pady=10, fill=tk.X)

        ttk.Label(input_frame, text="Keyword:").grid(row=0, column=0)
        self.keyword_entry = ttk.Entry(input_frame)
        self.keyword_entry.grid(row=0, column=1)

        ttk.Label(input_frame, text="Col Offset:").grid(row=0, column=2)
        self.col_offset_entry = ttk.Entry(input_frame, width=5)
        self.col_offset_entry.grid(row=0, column=3)

        ttk.Label(input_frame, text="Row Offset:").grid(row=0, column=4)
        self.row_offset_entry = ttk.Entry(input_frame, width=5)
        self.row_offset_entry.grid(row=0, column=5)

        ttk.Label(input_frame, text="New Value:").grid(row=0, column=6)
        self.new_value_entry = ttk.Entry(input_frame)
        self.new_value_entry.grid(row=0, column=7)

        ttk.Button(input_frame, text="Add Edit", command=self.add_edit).grid(row=0, column=8, padx=5)
        ttk.Button(input_frame, text="Delete Selected", command=self.delete_offset_edit).grid(row=0, column=9, padx=5)

        button_frame = ttk.Frame(self.master)
        button_frame.pack(pady=10, fill=tk.X)

        ttk.Button(button_frame, text="Clear Edits", command=self.clear_edits, padding=(12, 7, 12, 7)).pack(side=tk.RIGHT, padx=5, pady=5)
        ttk.Button(button_frame, text="Reset Confirmations", command=self.reset_confirmations, padding=(12, 7, 12, 7)).pack(side=tk.RIGHT, padx=5, pady=5)
        ttk.Button(button_frame, text="Select and Process Files", command=self.process_files, padding=(12, 7, 12, 7)).pack(side=tk.RIGHT, padx=5, pady=5)

    def import_edits(self):
        file_path = filedialog.askopenfilename(title="Import Edits", filetypes=[("JSON files", "*.json")])
        if file_path:
            try:
                with open(file_path, 'r') as f:
                    data = json.load(f)
                    
                # Handle both old and new format
                if isinstance(data, list):  # Old format: just a list of edits
                    self.edits = data
                    self.replacements = []
                    message = f"Imported {len(self.edits)} edits from old format file"
                else:  # New format: dictionary with edits and replacements
                    self.edits = data.get('edits', [])
                    self.replacements = []
                    for repl in data.get('replacements', []):
                        if len(repl) == 2:  # Convert old format to new
                            self.replacements.append((repl[0], repl[1], True))  # Default to exact match
                        else:
                            self.replacements.append(repl)
                    message = f"Imported {len(self.edits)} edits and {len(self.replacements)} replacements"
                    
                self.refresh_edits_tree()
                self.refresh_replace_tree()
                messagebox.showinfo("Import Successful", message)
            except Exception as e:
                messagebox.showerror("Import Error", f"Failed to import edits: {str(e)}")

    def export_edits(self):
        if not self.edits and not self.replacements:
            messagebox.showwarning("No Edits", "There are no edits or replacements to export.")
            return

        file_path = filedialog.asksaveasfilename(title="Export Edits", 
            defaultextension=".json", filetypes=[("JSON files", "*.json")])
        if file_path:
            try:
                data = {
                    'edits': self.edits,
                    'replacements': self.replacements
                }
                with open(file_path, 'w') as f:
                    json.dump(data, f, indent=2)
                messagebox.showinfo("Export Successful", 
                    f"Exported {len(self.edits)} edits and {len(self.replacements)} replacements to {file_path}")
            except Exception as e:
                messagebox.showerror("Export Error", f"Failed to export edits: {str(e)}")

    def refresh_edits_tree(self):
        self.edits_tree.delete(*self.edits_tree.get_children())
        for edit in self.edits:
            self.edits_tree.insert("", tk.END, values=edit)

    def refresh_replace_tree(self):
        self.replace_tree.delete(*self.replace_tree.get_children())
        for replacement in self.replacements:
            self.replace_tree.insert("", tk.END, values=replacement)

    def add_edit(self):
        keyword = self.keyword_entry.get()
        col_offset = self.col_offset_entry.get()
        row_offset = self.row_offset_entry.get()
        new_value = self.new_value_entry.get()

        if keyword and col_offset and row_offset and new_value:
            self.edits.append((keyword, int(col_offset), int(row_offset), new_value))
            self.edits_tree.insert("", tk.END, values=(keyword, col_offset, row_offset, new_value))
            self.clear_input_fields()
        else:
            messagebox.showwarning("Invalid Input", "Please fill all fields.")

    def clear_input_fields(self):
        self.keyword_entry.delete(0, tk.END)
        self.col_offset_entry.delete(0, tk.END)
        self.row_offset_entry.delete(0, tk.END)
        self.new_value_entry.delete(0, tk.END)

    def clear_edits(self):
        self.edits = []
        self.replacements = []  # Clear replacements too
        self.edits_tree.delete(*self.edits_tree.get_children())
        self.replace_tree.delete(*self.replace_tree.get_children())

    def reset_confirmations(self):
        """Reset confirmation settings to ask for each conflict"""
        self.confirmation_settings = {
            'replace_all': None,  # None = ask, True = always replace, False = always skip
            'skip_all': False
        }
        messagebox.showinfo("Confirmations Reset", "Confirmation settings have been reset. You will be asked for each conflict again.")

    def add_replace_pair(self):
        search_text = self.search_entry.get()
        replace_text = self.replace_entry.get()
        exact_match = self.exact_match_var.get()

        if search_text and replace_text:
            self.replacements.append((search_text, replace_text, exact_match))
            self.replace_tree.insert("", tk.END, values=(search_text, replace_text, "Yes" if exact_match else "No"))
            self.search_entry.delete(0, tk.END)
            self.replace_entry.delete(0, tk.END)
        else:
            messagebox.showwarning("Invalid Input", "Please fill both search and replace fields.")

    def delete_replace_pair(self):
        selected = self.replace_tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Please select a replacement pair to delete")
            return
            
        for item in selected:
            index = self.replace_tree.index(item)
            del self.replacements[index]
        self.refresh_replace_tree()

    def delete_offset_edit(self):
        selected = self.edits_tree.selection()
        if not selected:
            messagebox.showwarning("No Selection", "Please select an edit to delete")
            return
            
        for item in selected:
            index = self.edits_tree.index(item)
            del self.edits[index]
        self.refresh_edits_tree()

    def process_files(self):
        # Reset confirmation settings for new processing session
        self.confirmation_settings = {
            'replace_all': None,  # None = ask, True = always replace, False = always skip
            'skip_all': False
        }
        
        # Close existing output popup if it exists
        if self.output_popup and self.output_popup.winfo_exists():
            self.output_popup.destroy()
            
        self.show_output_popup()  # Create popup immediately
        # Start processing in a separate thread
        processing_thread = threading.Thread(target=self.run_processing)
        processing_thread.daemon = True
        processing_thread.start()

    def run_processing(self):
        try:
            selection = self.file_selection.get()
            if selection == "WALK":
                root_folder = filedialog.askdirectory(title='Select folder to walk through')
                if root_folder:
                    self.process_walk(root_folder)
            elif selection == "FOLDER":
                folder = filedialog.askdirectory(title='Select folder with Excel files')
                if folder:
                    self.process_folder(folder)
            elif selection == "FILES":
                file_paths = filedialog.askopenfilenames(
                    title="Select Excel files",
                    filetypes=[("Excel files", "*.xlsx")]
                )
                if file_paths:
                    for file_path in file_paths:
                        self.process_file(file_path)
            
            self.output_queue.put("Done Processing Files")
        except Exception as e:
            self.output_queue.put(f"Error during processing: {str(e)}")

    def show_output_popup(self):
        self.output_popup = tk.Toplevel(self.master)
        self.output_popup.title("Processing Output")
        
        text_frame = ttk.Frame(self.output_popup)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        text = tk.Text(text_frame, wrap=tk.WORD, height=20, width=80)
        text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        scrollbar = ttk.Scrollbar(text_frame, command=text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        text.config(yscrollcommand=scrollbar.set)
        text.insert(tk.END, "Processing started...\n")
        text.config(state=tk.DISABLED)
        
        close_button = ttk.Button(self.output_popup, text="Close", command=self.output_popup.destroy)
        close_button.pack(pady=5)

        def update_output():
            while not self.output_queue.empty():
                msg = self.output_queue.get()
                text.config(state=tk.NORMAL)
                text.insert(tk.END, msg + "\n")
                text.see(tk.END)
                text.config(state=tk.DISABLED)
            self.output_popup.after(100, update_output)  # Schedule next update
            
        self.output_popup.after(100, update_output)  # Start the update loop

    def process_walk(self, root_folder):
        for root, _, files in os.walk(root_folder):
            for file in files:
                if file.endswith('.xlsx') and not file.startswith('~'):
                    file_path = os.path.join(root, file)
                    self.make_edits(file_path)

    def process_folder(self, folder):
        for file in os.listdir(folder):
            if file.endswith('.xlsx') and not file.startswith('~'):
                file_path = os.path.join(folder, file)
                self.make_edits(file_path)

    def process_file(self, file_path):
        self.make_edits(file_path)

    def make_edits(self, file_path):
        try:
            app = xw.App(visible=False)
            wb = app.books.open(file_path)
            
            # Process keyword-based edits
            for edit in self.edits:
                keyword, col_offset, row_offset, new_value = edit
                self.edit_excel_cell(wb, keyword, col_offset, row_offset, new_value)
            
            # Process search/replace pairs
            for sheet in wb.sheets:
                used_range = sheet.used_range
                values = used_range.value
                if not values:  # Skip empty sheets
                    continue
                
                # Get the starting row and column of the used range
                start_row = used_range.row
                start_col = used_range.column
                    
                for i, row in enumerate(values):
                    for j, cell_value in enumerate(row):
                        if cell_value is not None:  # Only process non-None cells
                            str_value = str(cell_value)
                            for search_text, replace_text, exact_match in self.replacements:
                                if exact_match:
                                    if str_value == search_text:
                                        # Calculate the actual row and column in Excel's 1-based indexing
                                        actual_row = start_row + i
                                        actual_col = start_col + j
                                        
                                        target_cell = sheet.cells(actual_row, actual_col)
                                        existing_value = target_cell.value
                                        
                                        # Check if we should replace the value
                                        if self.should_replace_value(existing_value, replace_text, target_cell.address, sheet.name):
                                            target_cell.value = replace_text
                                            self.output_queue.put(f"REPLACED: '{existing_value}' → '{replace_text}' at {sheet.name}!{target_cell.address}")
                                        else:
                                            self.output_queue.put(f"SKIPPED: '{existing_value}' (no change needed) at {sheet.name}!{target_cell.address}")
                                else:
                                    if search_text in str_value:
                                        new_value = str_value.replace(search_text, replace_text)
                                        # Calculate the actual row and column in Excel's 1-based indexing
                                        actual_row = start_row + i
                                        actual_col = start_col + j
                                        
                                        target_cell = sheet.cells(actual_row, actual_col)
                                        existing_value = target_cell.value
                                        
                                        # Check if we should replace the value
                                        if self.should_replace_value(existing_value, new_value, target_cell.address, sheet.name):
                                            target_cell.value = new_value
                                            self.output_queue.put(f"REPLACED: '{existing_value}' → '{new_value}' at {sheet.name}!{target_cell.address}")
                                        else:
                                            self.output_queue.put(f"SKIPPED: '{existing_value}' (no change needed) at {sheet.name}!{target_cell.address}")
            
            wb.save()
            wb.close()
            app.quit()
            self.output_queue.put(f"Processed: {file_path}")
        except Exception as e:
            self.output_queue.put(f"Error processing {file_path}: {str(e)}")

    def edit_excel_cell(self, wb, keyword, col_offset, row_offset, new_value):
        for sheet in wb.sheets:
            used_range = sheet.used_range
            values = used_range.value
            if not values:  # Skip empty sheets
                continue
                
            # Find the starting row and column of the used range
            start_row = used_range.row
            start_col = used_range.column
            count = 0
            for i, row in enumerate(values):
                for j, cell_value in enumerate(row):
                    if cell_value == keyword:
                        # Calculate the actual row and column in Excel's 1-based indexing
                        actual_row = start_row + i + row_offset
                        actual_col = start_col + j + col_offset
                        
                        target_cell = sheet.cells(actual_row, actual_col)
                        existing_value = target_cell.value
                        
                        # Check if we should replace the value
                        if self.should_replace_value(existing_value, new_value, target_cell.address, sheet.name):
                            target_cell.value = new_value
                            self.output_queue.put(f"UPDATED: '{existing_value}' → '{new_value}' at {sheet.name}!{target_cell.address} (keyword '{keyword}' found at {sheet.cells(start_row + i, start_col + j).address})")
                            count += 1
                        else:
                            self.output_queue.put(f"SKIPPED: '{existing_value}' (no change needed) at {sheet.name}!{target_cell.address}")
            print(f"made {count} edits in {sheet.name}")
            
    def edit_replace_pair(self, event):
        item = self.replace_tree.selection()[0]
        values = self.replace_tree.item(item, 'values')
        
        edit_window = tk.Toplevel(self.master)
        edit_window.title("Edit Replace Pair")
        
        tk.Label(edit_window, text="Search Text:").grid(row=0, column=0)
        search_entry = ttk.Entry(edit_window, width=50)  # Make the entry field larger
        search_entry.grid(row=0, column=1)
        search_entry.insert(0, values[0])
        
        tk.Label(edit_window, text="Replace Text:").grid(row=1, column=0)
        replace_entry = ttk.Entry(edit_window, width=50)  # Make the entry field larger
        replace_entry.grid(row=1, column=1)
        replace_entry.insert(0, values[1])
        
        tk.Label(edit_window, text="Exact Match:").grid(row=2, column=0)
        exact_var = tk.BooleanVar(value=values[2] == "Yes" if len(values) > 2 else False)
        exact_check = ttk.Checkbutton(edit_window, text="", variable=exact_var)
        exact_check.grid(row=2, column=1, sticky=tk.W)
        
        def save_changes():
            new_search = search_entry.get()
            new_replace = replace_entry.get()
            if new_search and new_replace:
                index = self.replace_tree.index(item)
                self.replacements[index] = (new_search, new_replace, exact_var.get())
                self.replace_tree.item(item, values=(new_search, new_replace, "Yes" if exact_var.get() else "No"))
                edit_window.destroy()
            else:
                messagebox.showwarning("Invalid Input", "Both fields are required")
        
        ttk.Button(edit_window, text="Save", command=save_changes).grid(row=3, columnspan=2)
        
        # Center the window
        screen_width = self.master.winfo_screenwidth()
        screen_height = self.master.winfo_screenheight()
        x = (screen_width/2) - (edit_window.winfo_reqwidth()/2)
        y = (screen_height/2) - (edit_window.winfo_reqheight()/2)
        edit_window.geometry(f'+{int(x)}+{int(y)}')

    def edit_offset_edit(self, event):
        item = self.edits_tree.selection()[0]
        values = self.edits_tree.item(item, 'values')
        
        edit_window = tk.Toplevel(self.master)
        edit_window.title("Edit Offset Edit")
        edit_window.geometry("300x200")  # Set a fixed size for the window
        edit_window.resizable(False, False)  # Make the window non-resizable
        
        # Center the window
        screen_width = self.master.winfo_screenwidth()
        screen_height = self.master.winfo_screenheight()
        x = (screen_width/2) - (300/2)
        y = (screen_height/2) - (200/2)
        edit_window.geometry(f'300x200+{int(x)}+{int(y)}')
        
        fields = [
            ("Keyword:", 0, values[0]),
            ("Col Offset:", 1, values[1]),
            ("Row Offset:", 2, values[2]),
            ("New Value:", 3, values[3])
        ]
        
        entries = []
        for i, (label, col, value) in enumerate(fields):
            tk.Label(edit_window, text=label).grid(row=i, column=0)
            entry = ttk.Entry(edit_window, width=50)  # Make the entry field larger
            entry.grid(row=i, column=1)
            entry.insert(0, value)
            entries.append(entry)
        
        def save_changes():
            new_values = [e.get() for e in entries]
            if all(new_values[:3]) and new_values[3]:  # Check required fields
                try:
                    index = self.edits_tree.index(item)
                    self.edits[index] = (
                        new_values[0],
                        int(new_values[1]),
                        int(new_values[2]),
                        new_values[3]
                    )
                    self.refresh_edits_tree()
                    edit_window.destroy()
                except ValueError:
                    messagebox.showwarning("Invalid Input", "Offsets must be numbers")
            else:
                messagebox.showwarning("Invalid Input", "All fields are required")
        
        ttk.Button(edit_window, text="Save", command=save_changes).grid(row=4, columnspan=2)

if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelEditorApp(root)
    root.mainloop()
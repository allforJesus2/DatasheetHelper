import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog
import threading
import numpy as np
from sentence_transformers import SentenceTransformer
import openpyxl
from openpyxl.styles import Font, PatternFill
import os
from datetime import datetime


class SemanticMatcherApp:
    """
    A GUI application for semantic matching between two lists of text items.
    Uses sentence transformers to compute semantic similarity and exports results to Excel.
    """
    
    def __init__(self, root, parent_app=None):
        self.root = root
        self.parent_app = parent_app
        self.root.title("Semantic Matcher")
        self.root.geometry("900x700")
        
        # Initialize variables
        self.semantic_model = None
        self.model_loaded = False
        self.loading_model = False
        self.matching_in_progress = False
        self.results = []
        
        # Create the UI
        self.create_widgets()
        
        # Load the semantic model in background
        self.load_semantic_model_async()
    
    def create_widgets(self):
        """Create the main UI widgets"""
        # Main frame
        main_frame = ttk.Frame(self.root)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Title
        title_label = ttk.Label(main_frame, text="Semantic Matcher", 
                               font=("Arial", 16, "bold"))
        title_label.pack(pady=(0, 20))
        
        # Instructions
        instructions = ttk.Label(main_frame, 
                                text="Enter items in each list (one per line). The tool will match each item from List 1 to the most similar item in List 2.",
                                wraplength=800)
        instructions.pack(pady=(0, 10))
        
        # Input frame
        input_frame = ttk.Frame(main_frame)
        input_frame.pack(fill=tk.BOTH, expand=True)
        
        # List 1 frame
        list1_frame = ttk.LabelFrame(input_frame, text="List 1 (Source)", padding=10)
        list1_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 5))
        
        self.list1_text = scrolledtext.ScrolledText(list1_frame, height=15, wrap=tk.WORD)
        self.list1_text.pack(fill=tk.BOTH, expand=True)
        
        # List 2 frame
        list2_frame = ttk.LabelFrame(input_frame, text="List 2 (Target)", padding=10)
        list2_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(5, 0))
        
        self.list2_text = scrolledtext.ScrolledText(list2_frame, height=15, wrap=tk.WORD)
        self.list2_text.pack(fill=tk.BOTH, expand=True)
        
        # Controls frame
        controls_frame = ttk.Frame(main_frame)
        controls_frame.pack(fill=tk.X, pady=(10, 0))
        
        # Settings frame
        settings_frame = ttk.LabelFrame(controls_frame, text="Settings", padding=5)
        settings_frame.pack(fill=tk.X, pady=(0, 10))
        
        # Minimum similarity threshold
        threshold_frame = ttk.Frame(settings_frame)
        threshold_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(threshold_frame, text="Minimum Similarity Threshold:").pack(side=tk.LEFT)
        self.threshold_var = tk.DoubleVar(value=0.3)
        threshold_spinbox = ttk.Spinbox(threshold_frame, from_=0.0, to=1.0, increment=0.1, 
                                       textvariable=self.threshold_var, width=10)
        threshold_spinbox.pack(side=tk.LEFT, padx=(5, 0))
        
        ttk.Label(threshold_frame, text="(0.0 = match everything, 1.0 = exact matches only)").pack(side=tk.LEFT, padx=(10, 0))
        
        # Buttons frame
        buttons_frame = ttk.Frame(controls_frame)
        buttons_frame.pack(fill=tk.X, pady=5)
        
        # Match button
        self.match_button = ttk.Button(buttons_frame, text="Run Semantic Matching", 
                                      command=self.start_matching)
        self.match_button.pack(side=tk.LEFT, padx=(0, 5))
        
        # Clear button
        clear_button = ttk.Button(buttons_frame, text="Clear All", 
                                 command=self.clear_all)
        clear_button.pack(side=tk.LEFT, padx=5)
        
        # Export button
        self.export_button = ttk.Button(buttons_frame, text="Export to Excel", 
                                       command=self.export_to_excel, state='disabled')
        self.export_button.pack(side=tk.LEFT, padx=5)
        
        # Status and progress
        status_frame = ttk.Frame(controls_frame)
        status_frame.pack(fill=tk.X, pady=(5, 0))
        
        self.status_label = ttk.Label(status_frame, text="Loading semantic model...")
        self.status_label.pack(side=tk.LEFT)
        
        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(status_frame, variable=self.progress_var, 
                                           mode='determinate')
        self.progress_bar.pack(side=tk.RIGHT, fill=tk.X, expand=True, padx=(10, 0))
        
        # Results frame
        results_frame = ttk.LabelFrame(main_frame, text="Results", padding=10)
        results_frame.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        
        # Results treeview
        columns = ('source', 'matched_target', 'similarity', 'status')
        self.results_tree = ttk.Treeview(results_frame, columns=columns, show='headings', height=8)
        
        # Define headings
        self.results_tree.heading('source', text='Source Item')
        self.results_tree.heading('matched_target', text='Best Match')
        self.results_tree.heading('similarity', text='Similarity Score')
        self.results_tree.heading('status', text='Status')
        
        # Configure column widths
        self.results_tree.column('source', width=200)
        self.results_tree.column('matched_target', width=200)
        self.results_tree.column('similarity', width=100)
        self.results_tree.column('status', width=100)
        
        # Scrollbar for results
        results_scrollbar = ttk.Scrollbar(results_frame, orient=tk.VERTICAL, command=self.results_tree.yview)
        self.results_tree.configure(yscrollcommand=results_scrollbar.set)
        
        # Pack results tree and scrollbar
        self.results_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        results_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    
    def load_semantic_model_async(self):
        """Load the sentence transformer model in a background thread"""
        def load_model():
            try:
                self.loading_model = True
                self.status_label.config(text="Loading semantic model...")
                self.progress_bar.config(mode='indeterminate')
                self.progress_bar.start()
                
                # Use a lightweight model for faster loading
                self.semantic_model = SentenceTransformer('all-MiniLM-L6-v2')
                self.model_loaded = True
                
                # Update UI in main thread
                self.root.after(0, self.on_model_loaded)
                
            except Exception as e:
                # Update UI in main thread
                self.root.after(0, lambda: self.on_model_error(str(e)))
        
        thread = threading.Thread(target=load_model, daemon=True)
        thread.start()
    
    def on_model_loaded(self):
        """Called when the semantic model is successfully loaded"""
        self.loading_model = False
        self.progress_bar.stop()
        self.progress_bar.config(mode='determinate')
        self.progress_var.set(0)
        self.status_label.config(text="Ready - Semantic model loaded successfully")
        self.match_button.config(state='normal')
    
    def on_model_error(self, error_msg):
        """Called when there's an error loading the semantic model"""
        self.loading_model = False
        self.progress_bar.stop()
        self.progress_bar.config(mode='determinate')
        self.progress_var.set(0)
        self.status_label.config(text=f"Error loading model: {error_msg}")
        messagebox.showerror("Error", f"Failed to load semantic model: {error_msg}")
    
    def clear_all(self):
        """Clear all text inputs and results"""
        self.list1_text.delete(1.0, tk.END)
        self.list2_text.delete(1.0, tk.END)
        
        # Clear results
        for item in self.results_tree.get_children():
            self.results_tree.delete(item)
        
        self.results = []
        self.export_button.config(state='disabled')
        self.status_label.config(text="Ready - All data cleared")
    
    def start_matching(self):
        """Start the semantic matching process"""
        if not self.model_loaded or self.matching_in_progress:
            return
        
        # Get text from both inputs
        list1_text = self.list1_text.get(1.0, tk.END).strip()
        list2_text = self.list2_text.get(1.0, tk.END).strip()
        
        if not list1_text or not list2_text:
            messagebox.showwarning("Input Required", "Please enter items in both lists")
            return
        
        # Parse lists
        list1_items = [item.strip() for item in list1_text.split('\n') if item.strip()]
        list2_items = [item.strip() for item in list2_text.split('\n') if item.strip()]
        
        if not list1_items or not list2_items:
            messagebox.showwarning("Input Required", "Please enter at least one item in each list")
            return
        
        # Clear previous results
        for item in self.results_tree.get_children():
            self.results_tree.delete(item)
        
        # Start matching in background thread
        self.matching_in_progress = True
        self.match_button.config(state='disabled')
        self.export_button.config(state='disabled')
        
        def run_matching():
            try:
                self.root.after(0, lambda: self.status_label.config(text="Computing semantic similarities..."))
                
                results = []
                total_comparisons = len(list1_items)
                
                for i, source_item in enumerate(list1_items):
                    # Update progress
                    progress = (i + 1) / total_comparisons * 100
                    self.root.after(0, lambda p=progress: self.progress_var.set(p))
                    
                    best_match = None
                    best_score = -1
                    
                    # Compare with all items in list 2
                    for target_item in list2_items:
                        similarity = self.compute_semantic_similarity(source_item, target_item)
                        
                        if similarity > best_score:
                            best_score = similarity
                            best_match = target_item
                    
                    # Determine status based on threshold
                    threshold = self.threshold_var.get()
                    status = "Good Match" if best_score >= threshold else "Below Threshold"
                    
                    result = {
                        'source': source_item,
                        'matched_target': best_match,
                        'similarity': best_score,
                        'status': status
                    }
                    results.append(result)
                
                # Update UI in main thread
                self.root.after(0, lambda: self.on_matching_complete(results))
                
            except Exception as e:
                self.root.after(0, lambda: self.on_matching_error(str(e)))
        
        thread = threading.Thread(target=run_matching, daemon=True)
        thread.start()
    
    def compute_semantic_similarity(self, text1, text2):
        """Compute cosine similarity between two texts"""
        if not self.model_loaded or not self.semantic_model:
            return 0.0
        
        try:
            # Encode the texts
            embeddings = self.semantic_model.encode([text1, text2])
            
            # Calculate cosine similarity
            similarity = np.dot(embeddings[0], embeddings[1]) / (
                np.linalg.norm(embeddings[0]) * np.linalg.norm(embeddings[1])
            )
            
            return float(similarity)
        except Exception as e:
            print(f"Error computing semantic similarity: {e}")
            return 0.0
    
    def on_matching_complete(self, results):
        """Called when matching is complete"""
        self.matching_in_progress = False
        self.match_button.config(state='normal')
        self.progress_var.set(100)
        
        # Store results
        self.results = results
        
        # Display results in treeview
        for result in results:
            # Color code based on similarity score
            tags = []
            if result['similarity'] >= 0.7:
                tags = ['high_similarity']
            elif result['similarity'] >= 0.4:
                tags = ['medium_similarity']
            else:
                tags = ['low_similarity']
            
            self.results_tree.insert('', tk.END, 
                                   values=(result['source'], 
                                          result['matched_target'], 
                                          f"{result['similarity']:.3f}",
                                          result['status']),
                                   tags=tags)
        
        # Configure tags for color coding
        self.results_tree.tag_configure('high_similarity', background='lightgreen')
        self.results_tree.tag_configure('medium_similarity', background='lightyellow')
        self.results_tree.tag_configure('low_similarity', background='lightcoral')
        
        # Enable export button
        self.export_button.config(state='normal')
        
        # Update status
        good_matches = sum(1 for r in results if r['status'] == 'Good Match')
        self.status_label.config(text=f"Matching complete: {good_matches}/{len(results)} good matches")
    
    def on_matching_error(self, error_msg):
        """Called when there's an error during matching"""
        self.matching_in_progress = False
        self.match_button.config(state='normal')
        self.progress_var.set(0)
        self.status_label.config(text=f"Matching failed: {error_msg}")
        messagebox.showerror("Error", f"Semantic matching failed: {error_msg}")
    
    def export_to_excel(self):
        """Export results to Excel file"""
        if not self.results:
            messagebox.showwarning("No Results", "No matching results to export")
            return
        
        # Ask user for save location
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Save Semantic Matching Results"
        )
        
        if not filename:
            return
        
        try:
            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Semantic Matches"
            
            # Headers
            headers = ["Source Item", "Best Match", "Similarity Score", "Status", "Timestamp"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            # Data rows
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            for row, result in enumerate(self.results, 2):
                ws.cell(row=row, column=1, value=result['source'])
                ws.cell(row=row, column=2, value=result['matched_target'])
                ws.cell(row=row, column=3, value=result['similarity'])
                ws.cell(row=row, column=4, value=result['status'])
                ws.cell(row=row, column=5, value=timestamp)
                
                # Color code rows based on similarity
                if result['similarity'] >= 0.7:
                    fill_color = "C6EFCE"  # Light green
                elif result['similarity'] >= 0.4:
                    fill_color = "FFEB9C"  # Light yellow
                else:
                    fill_color = "FFC7CE"  # Light red
                
                for col in range(1, 6):
                    ws.cell(row=row, column=col).fill = PatternFill(
                        start_color=fill_color, end_color=fill_color, fill_type="solid"
                    )
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Add summary sheet
            summary_ws = wb.create_sheet("Summary")
            summary_data = [
                ["Semantic Matching Summary", ""],
                ["Generated:", timestamp],
                ["Total Source Items:", len(self.results)],
                ["Threshold Used:", self.threshold_var.get()],
                ["Good Matches:", sum(1 for r in self.results if r['status'] == 'Good Match')],
                ["Below Threshold:", sum(1 for r in self.results if r['status'] == 'Below Threshold')],
                ["Average Similarity:", f"{np.mean([r['similarity'] for r in self.results]):.3f}"],
                ["Best Similarity:", f"{max(r['similarity'] for r in self.results):.3f}"],
                ["Worst Similarity:", f"{min(r['similarity'] for r in self.results):.3f}"]
            ]
            
            for row, (label, value) in enumerate(summary_data, 1):
                summary_ws.cell(row=row, column=1, value=label).font = Font(bold=True)
                summary_ws.cell(row=row, column=2, value=value)
            
            # Save file
            wb.save(filename)
            
            messagebox.showinfo("Export Complete", f"Results exported successfully to:\n{filename}")
            self.status_label.config(text=f"Results exported to {os.path.basename(filename)}")
            
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export results: {str(e)}")


if __name__ == "__main__":
    # Test the application standalone
    root = tk.Tk()
    app = SemanticMatcherApp(root)
    root.mainloop()

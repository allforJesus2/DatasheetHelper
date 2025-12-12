import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import os
import re
import subprocess
import threading
import time
from collections import deque
from io import BytesIO
import pdfplumber
from openpyxl import Workbook

# Import libraries for OCR fallback
try:
    import fitz  # PyMuPDF
    fitz_available = True
except ImportError:
    fitz_available = False

try:
    from PIL import Image
    import numpy as np
    pil_available = True
    numpy_available = True
except ImportError:
    pil_available = False
    numpy_available = False

try:
    import easyocr
    easyocr_available = True
    # Initialize EasyOCR reader (lazy loading)
    easyocr_reader = None
except ImportError:
    easyocr_available = False
    easyocr_reader = None

MAX_HISTORY = 15  # Max number of regex patterns to store

class PDFRegexSearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("PDF Regex Search")
        self.root.geometry("800x600")
        
        # Variables
        self.folder_path = tk.StringVar()
        self.selected_files = []  # List of selected file paths
        self.regex_pattern = tk.StringVar()
        self.results = []
        self.is_searching = False
        self.should_cancel = False  # Add flag for cancellation
        self.last_status_message = "Ready"  # Store the last search status message
        self.filtered_items = []  # Store filtered items
        self.current_filter = {}  # Store active filters by column
        
        # Sorting variables
        self.sort_column = None
        self.sort_reverse = False
        
        # History setup
        self.regex_history = deque(maxlen=MAX_HISTORY)
        self.history_file = os.path.join(os.path.expanduser("~"), ".pdfregexsearch_history.txt")
        self.load_history()
        
        # Create UI
        self.create_widgets()
        
        # Save history on close
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        
    def load_history(self):
        try:
            if os.path.exists(self.history_file):
                with open(self.history_file, 'r') as f:
                    # Read lines, strip whitespace, filter empty lines
                    history = [line.strip() for line in f if line.strip()]
                    # Use deque's extend which respects maxlen
                    self.regex_history.extend(history) 
        except Exception as e:
            print(f"Error loading regex history: {e}")

    def save_history(self):
        try:
            with open(self.history_file, 'w') as f:
                # Write history, one pattern per line (most recent first)
                for pattern in self.regex_history:
                    f.write(pattern + '\n')
        except Exception as e:
            print(f"Error saving regex history: {e}")

    def update_regex_history(self, pattern):
        if not pattern:  # Don't add empty patterns
            return
            
        # Remove if exists to avoid duplicates and move to front
        if pattern in self.regex_history:
            self.regex_history.remove(pattern)
            
        # Add to the left (most recent)
        self.regex_history.appendleft(pattern)
        
        # Update combobox values (if it exists)
        if hasattr(self, 'regex_combobox'):
            self.regex_combobox['values'] = list(self.regex_history)

    def on_close(self):
        self.save_history()
        self.root.destroy()

    def create_widgets(self):
        # File/Folder selection
        selection_frame = ttk.LabelFrame(self.root, text="Search Scope", padding=10)
        selection_frame.pack(fill=tk.X, padx=10, pady=5)
        
        # Selection mode
        mode_frame = ttk.Frame(selection_frame)
        mode_frame.pack(fill=tk.X, pady=2)
        
        self.search_mode = tk.StringVar(value="folder")
        ttk.Radiobutton(mode_frame, text="Search Folder", variable=self.search_mode, 
                       value="folder", command=self.on_mode_change).pack(side=tk.LEFT)
        ttk.Radiobutton(mode_frame, text="Select Files", variable=self.search_mode, 
                       value="files", command=self.on_mode_change).pack(side=tk.LEFT, padx=10)
        
        # Folder selection
        folder_frame = ttk.Frame(selection_frame)
        folder_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(folder_frame, text="Folder:").pack(side=tk.LEFT)
        self.folder_entry = ttk.Entry(folder_frame, textvariable=self.folder_path, width=50)
        self.folder_entry.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        # Bind Enter key to run search
        self.folder_entry.bind("<Return>", lambda event: self.start_search())
        ttk.Button(folder_frame, text="Browse...", command=self.browse_folder).pack(side=tk.LEFT)
        
        # File selection
        files_frame = ttk.Frame(selection_frame)
        files_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(files_frame, text="Files:").pack(side=tk.LEFT)
        self.files_label = ttk.Label(files_frame, text="No files selected", foreground="gray")
        self.files_label.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(files_frame, text="Select Files...", command=self.select_files).pack(side=tk.LEFT)
        ttk.Button(files_frame, text="Clear", command=self.clear_selected_files).pack(side=tk.LEFT, padx=5)
        
        # Initially hide file selection
        files_frame.pack_forget()
        
        # Regex pattern - Use Combobox
        pattern_frame = ttk.Frame(self.root, padding=10)
        pattern_frame.pack(fill=tk.X)
        
        ttk.Label(pattern_frame, text="Regex Pattern:").pack(side=tk.LEFT)
        # Changed Entry to Combobox
        self.regex_combobox = ttk.Combobox(pattern_frame, textvariable=self.regex_pattern, width=48) 
        self.regex_combobox.pack(side=tk.LEFT, padx=5)
        # Bind Enter key to run search
        self.regex_combobox.bind("<Return>", lambda event: self.start_search())
        
        # Populate with loaded history
        self.regex_combobox['values'] = list(self.regex_history) 
        # Set current value if history exists
        if self.regex_history:
            self.regex_pattern.set(self.regex_history[0])
        
        # Bind pattern changes to auto-test
        self.regex_pattern.trace('w', self.on_test_text_change) 

        self.search_button = ttk.Button(pattern_frame, text="Search", command=self.start_search)
        self.search_button.pack(side=tk.LEFT)
        
        # Add Test Regex button
        self.test_button = ttk.Button(pattern_frame, text="Test Regex", command=self.test_regex)
        self.test_button.pack(side=tk.LEFT, padx=5)
        
        # Add cancel button
        self.cancel_button = ttk.Button(pattern_frame, text="Cancel", command=self.cancel_search, state='disabled')
        self.cancel_button.pack(side=tk.LEFT, padx=5)
        
        # Export button
        self.export_button = ttk.Button(pattern_frame, text="Export Results", command=self.export_results, state='disabled')
        self.export_button.pack(side=tk.LEFT, padx=5)
        
        # Regex testing section (collapsible)
        self.test_frame = ttk.LabelFrame(self.root, text="Test Regex Pattern", padding=5)
        self.test_frame.pack(fill=tk.X, padx=10, pady=5)
        
        # Test text input (small)
        test_input_frame = ttk.Frame(self.test_frame)
        test_input_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(test_input_frame, text="Test Text:").pack(side=tk.LEFT)
        self.test_text_var = tk.StringVar()
        self.test_text_entry = ttk.Entry(test_input_frame, textvariable=self.test_text_var, width=50)
        self.test_text_entry.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        # Bind text changes to auto-test
        self.test_text_var.trace('w', self.on_test_text_change)
        
        # Sample text button
        ttk.Button(test_input_frame, text="Sample", 
                  command=self.load_sample_text).pack(side=tk.LEFT, padx=5)
        
        # Test results (small listbox)
        test_results_frame = ttk.Frame(self.test_frame)
        test_results_frame.pack(fill=tk.X, pady=2)
        
        ttk.Label(test_results_frame, text="Matches:").pack(side=tk.LEFT)
        
        # Create a small listbox for test results
        self.test_results_listbox = tk.Listbox(test_results_frame, height=3)
        test_results_scrollbar = ttk.Scrollbar(test_results_frame, orient="vertical", command=self.test_results_listbox.yview)
        self.test_results_listbox.configure(yscrollcommand=test_results_scrollbar.set)
        
        self.test_results_listbox.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        test_results_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Initially hide the test frame
        self.test_frame.pack_forget()
        
        # Status bar
        self.status_var = tk.StringVar()
        self.status_var.set("Ready")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Progress bar
        self.progress_var = tk.DoubleVar()
        self.progress = ttk.Progressbar(self.root, variable=self.progress_var, maximum=100)
        self.progress.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Results display
        results_frame = ttk.LabelFrame(self.root, text="Results", padding=10)
        results_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Create column header context menu 
        self.header_menu = tk.Menu(self.root, tearoff=0)
        self.header_menu.add_command(label="Filter Column...", command=self.filter_column)
        self.header_menu.add_command(label="Clear Column Filter", command=self.clear_column_filter)
        self.header_menu.add_separator()
        self.header_menu.add_command(label="Copy Column to Clipboard", command=self.copy_column_to_clipboard)
        self.header_menu.add_separator()
        self.header_menu.add_command(label="Clear All Filters", command=self.clear_all_filters)
        
        # Create treeview with scrollbars
        self.tree = ttk.Treeview(results_frame, columns=("File", "Page", "Match"), show="headings")
        
        # Configure column headings with sorting click handlers
        for col in self.tree["columns"]:
            self.tree.heading(col, text=col, command=lambda _col=col: self.sort_column_data(_col))
        
        self.tree.column("File", width=200)
        self.tree.column("Page", width=80)
        self.tree.column("Match", width=400)
        
        # Bind right-click on column headers
        self.tree.bind("<ButtonPress-3>", self.on_header_right_click)
        
        # Store column index for filtering
        self.filter_column_index = None
        
        # Scrollbars
        vsb = ttk.Scrollbar(results_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(results_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        
        # Grid layout for treeview and scrollbars
        self.tree.grid(column=0, row=0, sticky='nsew')
        vsb.grid(column=1, row=0, sticky='ns')
        hsb.grid(column=0, row=1, sticky='ew')
        
        results_frame.grid_columnconfigure(0, weight=1)
        results_frame.grid_rowconfigure(0, weight=1)
        
        # Bind double-click event
        self.tree.bind("<Double-1>", self.open_pdf_file)
        
    def test_regex(self):
        """Toggle regex testing section"""
        if self.test_frame.winfo_ismapped():
            # Hide the test frame
            self.test_frame.pack_forget()
        else:
            # Show the test frame
            self.test_frame.pack(fill=tk.X, padx=10, pady=5, after=self.regex_combobox.master)
            # Test the current pattern if there's test text
            if self.test_text_var.get():
                self.run_test()
    
    def load_sample_text(self):
        """Load sample text for testing"""
        sample_texts = [
            "123-456-7890, ABC123, test@email.com, 2023-12-25",
            "PRD-001, PRD-002, PRD-003",
            "(555) 123-4567, 555-987-6543, +1-555-123-4567",
            "john.doe@company.com, jane.smith@test.org",
            "2023/12/25, 12-25-2023, 25 Dec 2023"
        ]
        
        # Combine into one line for the entry widget
        combined_text = " | ".join(sample_texts)
        self.test_text_var.set(combined_text)
        
        # Run test if there's a pattern
        if self.regex_pattern.get():
            self.run_test()
    
    def run_test(self):
        """Run the regex test with current pattern and test text"""
        pattern = self.regex_pattern.get()
        test_text = self.test_text_var.get()
        
        if not pattern:
            return
            
        if not test_text:
            return
            
        # Clear previous results
        self.test_results_listbox.delete(0, tk.END)
        
        # Test the pattern
        results, error = self.test_regex_pattern(pattern, test_text)
        
        if error:
            self.test_results_listbox.insert(tk.END, f"Error: {error}")
            return
            
        if results:
            for result in results:
                display_text = f"'{result['match']}' (pos {result['start']}-{result['end']})"
                if result['groups'] != 'None':
                    display_text += f" [groups: {result['groups']}]"
                self.test_results_listbox.insert(tk.END, display_text)
        else:
            self.test_results_listbox.insert(tk.END, "No matches found")
    
    def on_test_text_change(self, *args):
        """Called when test text changes - run test automatically"""
        if self.test_frame.winfo_ismapped() and self.regex_pattern.get():
            self.run_test()
    
    def test_regex_pattern(self, pattern, test_text):
        """Test a regex pattern against test text and return results"""
        try:
            regex = re.compile(pattern)
            matches = list(regex.finditer(test_text))
            
            results = []
            for match in matches:
                groups = match.groups()
                groups_str = ", ".join([f"'{g}'" for g in groups]) if groups else "None"
                
                results.append({
                    "match": match.group(),
                    "start": match.start(),
                    "end": match.end(),
                    "groups": groups_str
                })
            
            return results, None
        except re.error as e:
            return None, str(e)
    
    def on_mode_change(self):
        """Handle search mode change"""
        if self.search_mode.get() == "folder":
            # Show folder selection, hide file selection
            self.folder_entry.master.pack(fill=tk.X, pady=2)
            self.files_label.master.pack_forget()
        else:
            # Show file selection, hide folder selection
            self.folder_entry.master.pack_forget()
            self.files_label.master.pack(fill=tk.X, pady=2)
    
    def browse_folder(self):
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.folder_path.set(folder_selected)
    
    def select_files(self):
        """Open file dialog to select PDF files"""
        files_selected = filedialog.askopenfilenames(
            title="Select PDF Files",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")]
        )
        if files_selected:
            self.selected_files = list(files_selected)
            self.update_files_label()
    
    def clear_selected_files(self):
        """Clear selected files"""
        self.selected_files = []
        self.update_files_label()
    
    def update_files_label(self):
        """Update the files label to show selected files"""
        if not self.selected_files:
            self.files_label.config(text="No files selected", foreground="gray")
        elif len(self.selected_files) == 1:
            filename = os.path.basename(self.selected_files[0])
            self.files_label.config(text=f"1 file: {filename}", foreground="black")
        else:
            filenames = [os.path.basename(f) for f in self.selected_files[:3]]
            if len(self.selected_files) > 3:
                filenames.append(f"... and {len(self.selected_files) - 3} more")
            self.files_label.config(text=f"{len(self.selected_files)} files: {', '.join(filenames)}", foreground="black")
    
    def start_search(self):
        if self.is_searching:
            return
            
        pattern = self.regex_pattern.get()  # Get pattern from the combobox's variable
        
        if not pattern:
            messagebox.showerror("Error", "Please enter or select a regex pattern")
            return
            
        try:
            re.compile(pattern)
        except re.error:
            messagebox.showerror("Error", "Invalid regular expression")
            return
            
        # Validate search scope based on mode
        if self.search_mode.get() == "folder":
            folder = self.folder_path.get()
            if not folder or not os.path.isdir(folder):
                messagebox.showerror("Error", "Please select a valid folder")
                return
            files_to_search = None  # Will search all PDF files in folder
        else:
            if not self.selected_files:
                messagebox.showerror("Error", "Please select at least one PDF file")
                return
            files_to_search = self.selected_files
            
        # --- Add pattern to history ---
        self.update_regex_history(pattern)
        # --- End Add pattern ---
            
        # Clear previous results and reset filters/sort
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        self.current_filter.clear()
        self.sort_column = None
        self.sort_reverse = False
        
        # Reset column headings to remove sort indicators
        for col in self.tree["columns"]:
            self.tree.heading(col, text=col)
        
        self.results = []
        self.is_searching = True
        self.should_cancel = False
        self.search_button.state(['disabled'])
        self.cancel_button.state(['!disabled'])  # Enable cancel button
        self.export_button.state(['disabled'])
        self.status_var.set("Searching...")
        self.progress_var.set(0)
        
        # Start search in a separate thread
        if self.search_mode.get() == "folder":
            thread = threading.Thread(target=self.search_pdf_files, args=(folder, pattern))
        else:
            thread = threading.Thread(target=self.search_selected_files, args=(files_to_search, pattern))
        thread.daemon = True
        thread.start()
    
    def cancel_search(self):
        if self.is_searching:
            self.should_cancel = True
            self.status_var.set("Cancelling search...")
    
    def extract_text_with_ocr(self, pdf_path):
        """Extract text from PDF using OCR when no text is found
        Returns a list of tuples: [(page_num, text), ...]
        """
        global easyocr_reader
        
        if not fitz_available or not pil_available or not numpy_available or not easyocr_available:
            return []
        
        # Initialize EasyOCR reader if not already done
        if easyocr_reader is None:
            easyocr_reader = easyocr.Reader(['en'], gpu=False)
        
        # Convert PDF to images using fitz
        try:
            pdf_document = fitz.open(pdf_path)
        except Exception:
            return []
        
        page_texts = []
        try:
            for page_num in range(len(pdf_document)):
                if self.should_cancel:
                    break
                
                # Get the page
                page = pdf_document[page_num]
                
                # Render page to a pixmap (image)
                pixmap = page.get_pixmap()
                
                # Convert pixmap to numpy array for EasyOCR
                img_data = pixmap.tobytes("ppm")
                image = Image.open(BytesIO(img_data))
                img_array = np.array(image)
                
                # Perform OCR on the numpy array
                results = easyocr_reader.readtext(img_array)
                # Combine all text from OCR results
                page_text = "\n".join([result[1] for result in results])
                if page_text:
                    page_texts.append((page_num + 1, page_text))
        finally:
            pdf_document.close()
        
        return page_texts
    
    def extract_text_from_pdf(self, file_path, reg, file_name):
        """Extract text from PDF, using OCR if no text is found"""
        text_extracted = ""
        pages_with_text = 0
        
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    if self.should_cancel:
                        break
                    
                    text = page.extract_text()
                    if text:
                        text_extracted += text
                        pages_with_text += 1
                        # Search for matches in the text
                        matches = list(reg.finditer(text))
                        for match in matches:
                            result = {
                                "file": file_name,
                                "file_path": file_path,
                                "page": page_num,
                                "match": match.group()
                            }
                            self.results.append(result)
                            # Add to treeview
                            self.root.after(0, lambda r=result: self.add_result_to_tree(r))
        except Exception as e:
            print(f"Error processing {file_path} with pdfplumber: {str(e)}")
        
        # If no text was found, try OCR
        if not text_extracted and fitz_available and pil_available and numpy_available and easyocr_available:
            try:
                self.root.after(0, lambda f=file_name: self.status_var.set(f"Using OCR on {f}..."))
                ocr_pages = self.extract_text_with_ocr(file_path)
                if ocr_pages:
                    # Search for matches in OCR text per page
                    for page_num, page_text in ocr_pages:
                        if self.should_cancel:
                            break
                        matches = list(reg.finditer(page_text))
                        for match in matches:
                            result = {
                                "file": file_name,
                                "file_path": file_path,
                                "page": page_num,
                                "match": match.group()
                            }
                            self.results.append(result)
                            # Add to treeview
                            self.root.after(0, lambda r=result: self.add_result_to_tree(r))
        elif not text_extracted:
            # No text found and OCR not available
            print(f"No text found in {file_name} and OCR not available")
    
    def search_pdf_files(self, folder, pattern):
        try:
            # Get all PDF files in the folder
            pdf_files = [f for f in os.listdir(folder) if f.endswith('.pdf')]
            total_files = len(pdf_files)
            
            if total_files == 0:
                self.root.after(0, lambda: self.status_var.set("No PDF files found in the folder"))
                self.root.after(0, lambda: self.search_button.state(['!disabled']))
                self.root.after(0, lambda: self.cancel_button.state(['disabled']))
                self.is_searching = False
                return
                
            reg = re.compile(pattern)
            
            for i, file in enumerate(pdf_files):
                if self.should_cancel:
                    self.root.after(0, lambda: self.status_var.set("Search cancelled"))
                    break
                    
                file_path = os.path.join(folder, file)
                try:
                    # Update progress
                    progress = (i / total_files) * 100
                    self.root.after(0, lambda p=progress: self.progress_var.set(p))
                    self.root.after(0, lambda f=file: self.status_var.set(f"Searching in {f}..."))
                    
                    # Extract text from PDF and search for matches
                    self.extract_text_from_pdf(file_path, reg, file)
                    
                except Exception as e:
                    print(f"Error processing {file}: {str(e)}")
            
            # Update final status
            if not self.should_cancel:
                self.root.after(0, lambda: self.progress_var.set(100))
                final_message = f"Search complete. Found {len(self.results)} matches."
                self.root.after(0, lambda msg=final_message: self.status_var.set(msg))
                self.root.after(0, lambda msg=final_message: setattr(self, 'last_status_message', msg))
            else:
                # Update status if cancelled
                cancel_message = "Search cancelled"
                self.root.after(0, lambda msg=cancel_message: self.status_var.set(msg))
                self.root.after(0, lambda msg=cancel_message: setattr(self, 'last_status_message', msg))
                
            # Enable export button if results were found
            if self.results and not self.should_cancel:  # Only enable if search completed and found results
                self.root.after(0, lambda: self.export_button.state(['!disabled']))
            
        except Exception as e:
            error_message = f"Search error: {str(e)}"
            self.root.after(0, lambda: messagebox.showerror("Error", error_message))
            self.root.after(0, lambda msg=error_message: self.status_var.set(msg))  # Show error in status too
            self.root.after(0, lambda msg=error_message: setattr(self, 'last_status_message', msg))
        finally:
            self.root.after(0, lambda: self.search_button.state(['!disabled']))
            self.root.after(0, lambda: self.cancel_button.state(['disabled']))
            self.is_searching = False
    
    def search_selected_files(self, file_paths, pattern):
        """Search in specific selected files"""
        try:
            total_files = len(file_paths)
            reg = re.compile(pattern)
            
            for i, file_path in enumerate(file_paths):
                if self.should_cancel:
                    self.root.after(0, lambda: self.status_var.set("Search cancelled"))
                    break
                    
                try:
                    # Update progress
                    progress = (i / total_files) * 100
                    self.root.after(0, lambda p=progress: self.progress_var.set(p))
                    filename = os.path.basename(file_path)
                    self.root.after(0, lambda f=filename: self.status_var.set(f"Searching in {f}..."))
                    
                    # Extract text from PDF and search for matches
                    self.extract_text_from_pdf(file_path, reg, filename)
                    
                except Exception as e:
                    print(f"Error processing {file_path}: {str(e)}")
            
            # Update final status
            if not self.should_cancel:
                self.root.after(0, lambda: self.progress_var.set(100))
                final_message = f"Search complete. Found {len(self.results)} matches."
                self.root.after(0, lambda msg=final_message: self.status_var.set(msg))
                self.root.after(0, lambda msg=final_message: setattr(self, 'last_status_message', msg))
            else:
                # Update status if cancelled
                cancel_message = "Search cancelled"
                self.root.after(0, lambda msg=cancel_message: self.status_var.set(msg))
                self.root.after(0, lambda msg=cancel_message: setattr(self, 'last_status_message', msg))
                
            # Enable export button if results were found
            if self.results and not self.should_cancel:  # Only enable if search completed and found results
                self.root.after(0, lambda: self.export_button.state(['!disabled']))
            
        except Exception as e:
            error_message = f"Search error: {str(e)}"
            self.root.after(0, lambda: messagebox.showerror("Error", error_message))
            self.root.after(0, lambda msg=error_message: self.status_var.set(msg))  # Show error in status too
            self.root.after(0, lambda msg=error_message: setattr(self, 'last_status_message', msg))
        finally:
            self.root.after(0, lambda: self.search_button.state(['!disabled']))
            self.root.after(0, lambda: self.cancel_button.state(['disabled']))
            self.is_searching = False
    
    def add_result_to_tree(self, result):
        self.tree.insert("", "end", values=(
            result["file"],
            result["page"],
            result["match"]
        ))
    
    def open_pdf_file(self, event):
        # Get selected item
        selected_item = self.tree.focus()
        if not selected_item:
            return
            
        # Get the corresponding result values
        values = self.tree.item(selected_item, "values")
        if not values or len(values) < 3:
            return  # Should not happen, but good practice
            
        file_name = values[0]
        page_num = values[1]
        
        # Find the full file path from the stored results
        file_path = None
        for result in self.results:
            if result["file"] == file_name and result["page"] == int(page_num):
                file_path = result["file_path"]
                break
        
        if not file_path:
            messagebox.showerror("Error", "Could not find file path for the selected result.")
            return

        # Store current status and show loading message
        previous_status = self.status_var.get()
        self.status_var.set(f"Opening {os.path.basename(file_path)}...")
        self.root.update_idletasks()  # Force UI update
        
        try:
            # Open PDF with default application
            if os.name == 'nt':  # Windows
                os.startfile(file_path)
            elif os.name == 'posix':  # macOS and Linux
                if os.uname().sysname == 'Darwin':  # macOS
                    subprocess.call(('open', file_path))
                else:  # Linux
                    subprocess.call(('xdg-open', file_path))
        except Exception as e:
            messagebox.showerror("Error", f"Could not open PDF file: {str(e)}")
        finally:
            # Restore previous status message
            self.status_var.set(self.last_status_message)  # Restore the last search status

    def export_results(self):
        if not self.results:
            messagebox.showinfo("Export", "No results to export")
            return
            
        # Ask user for save location
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Save Results As"
        )
        
        if not file_path:
            return  # User cancelled
            
        try:
            # Create a new workbook and select the active worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Search Results"
            
            # Add headers
            headers = ["File", "Page", "Match"]
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            
            # Add data
            for row_idx, result in enumerate(self.results, 2):
                ws.cell(row=row_idx, column=1, value=result["file"])
                ws.cell(row=row_idx, column=2, value=result["page"])
                ws.cell(row=row_idx, column=3, value=result["match"])
            
            # Auto-adjust column widths
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column].width = adjusted_width
            
            # Save the workbook
            wb.save(file_path)
            
            # Ask if user wants to open the exported file
            if messagebox.askyesno("Export Successful", f"Results exported to {file_path}\n\nWould you like to open the file?"):
                try:
                    # Open the Excel file with default application
                    if os.name == 'nt':  # Windows
                        os.startfile(file_path)
                    elif os.name == 'posix':  # macOS and Linux
                        if os.uname().sysname == 'Darwin':  # macOS
                            subprocess.call(('open', file_path))
                        else:  # Linux
                            subprocess.call(('xdg-open', file_path))
                except Exception as e:
                    messagebox.showerror("Error", f"Could not open file: {str(e)}")
            
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export results: {str(e)}")

    def on_header_right_click(self, event):
        # Get the column that was right-clicked
        region = self.tree.identify_region(event.x, event.y)
        if region == "heading":
            column_id = self.tree.identify_column(event.x)
            column_index = int(column_id.replace('#', '')) - 1
            column_name = self.tree["columns"][column_index]
            
            # Store the column being filtered
            self.filter_column_index = column_index
            
            # Update menu labels with column name
            if column_name in self.current_filter:
                filter_text = f"Filter '{column_name}' (active)..."
                clear_text = f"Clear '{column_name}' Filter"
            else:
                filter_text = f"Filter '{column_name}'..."
                clear_text = f"Clear '{column_name}' Filter"
                
            self.header_menu.entryconfig(0, label=filter_text)
            self.header_menu.entryconfig(1, label=clear_text)
            self.header_menu.entryconfig(3, label=f"Copy '{column_name}' to Clipboard")
            
            # Show menu at event position
            self.header_menu.post(event.x_root, event.y_root)
            
    def filter_column(self):
        if self.filter_column_index is None:
            return
        
        column_name = self.tree["columns"][self.filter_column_index]
        
        # Get filter string from user
        filter_string = simpledialog.askstring(
            "Filter Column", 
            f"Enter text to filter in '{column_name}' column:",
            initialvalue=self.current_filter.get(column_name, "")
        )
        
        if filter_string is None:  # User canceled
            return
            
        if filter_string.strip() == "":
            self.clear_column_filter()
            return
            
        # Apply the filter
        self.current_filter[column_name] = filter_string
        self.apply_filters()
        
        # Update status bar
        self.status_var.set(f"Filtered by {column_name}: '{filter_string}'")
        
    def clear_column_filter(self):
        if self.filter_column_index is None:
            return
            
        column_name = self.tree["columns"][self.filter_column_index]
        
        if column_name in self.current_filter:
            del self.current_filter[column_name]
            self.apply_filters()
            
            if self.current_filter:
                self.status_var.set(f"Cleared filter for {column_name}, other filters still active")
            else:
                self.status_var.set("All filters cleared")
                
    def clear_all_filters(self):
        if not self.current_filter:
            return
            
        self.current_filter.clear()
        self.apply_filters()
        self.status_var.set("All filters cleared")
    
    def copy_column_to_clipboard(self):
        if self.filter_column_index is None:
            return
            
        column_name = self.tree["columns"][self.filter_column_index]
        
        # Get all values for the selected column from filtered results
        column_values = []
        
        # Get the filtered results (same logic as apply_filters)
        if not self.current_filter:
            filtered_results = self.results.copy()
        else:
            filtered_results = []
            for result in self.results:
                matches_all_filters = True
                
                for col_name, filter_text in self.current_filter.items():
                    result_value = ""
                    
                    if col_name == "File":
                        result_value = result["file"]
                    elif col_name == "Page":
                        result_value = str(result["page"])
                    elif col_name == "Match":
                        result_value = result["match"]
                        
                    if filter_text.lower() not in result_value.lower():
                        matches_all_filters = False
                        break
                        
                if matches_all_filters:
                    filtered_results.append(result)
        
        # Extract values for the selected column
        for result in filtered_results:
            if column_name == "File":
                column_values.append(result["file"])
            elif column_name == "Page":
                column_values.append(str(result["page"]))
            elif column_name == "Match":
                column_values.append(result["match"])
        
        if column_values:
            # Join values with newlines
            clipboard_text = "\n".join(column_values)
            
            # Copy to clipboard
            self.root.clipboard_clear()
            self.root.clipboard_append(clipboard_text)
            
            # Update status
            self.status_var.set(f"Copied {len(column_values)} {column_name.lower()} values to clipboard")
        else:
            self.status_var.set(f"No {column_name.lower()} values to copy")
        
    def sort_column_data(self, column):
        """Sort tree contents when a column header is clicked"""
        if self.sort_column == column:
            # If already sorting by this column, reverse the sort order
            self.sort_reverse = not self.sort_reverse
        else:
            # First time sorting by this column
            self.sort_column = column
            self.sort_reverse = False
            
        # Update column headings to show sort indicators
        for col in self.tree["columns"]:
            if col == self.sort_column:
                direction = " ▼" if self.sort_reverse else " ▲"
                self.tree.heading(col, text=f"{col}{direction}")
            else:
                self.tree.heading(col, text=col)
                
        # Apply the sort to the displayed items
        self.apply_filters()  # This will reapply filters and also sort the data
        
        # Update status bar
        direction = "descending" if self.sort_reverse else "ascending"
        self.status_var.set(f"Sorted by {column} ({direction})")
    
    def apply_filters(self):
        # Clear treeview
        for item in self.tree.get_children():
            self.tree.delete(item)
            
        # Get the filtered results
        filtered_results = []
        
        # With no filters, use all results
        if not self.current_filter:
            filtered_results = self.results.copy()
        else:
            # Apply filters
            for result in self.results:
                # Check all active filters
                matches_all_filters = True
                
                for column_name, filter_text in self.current_filter.items():
                    result_value = ""
                    
                    if column_name == "File":
                        result_value = result["file"]
                    elif column_name == "Page":
                        result_value = str(result["page"])
                    elif column_name == "Match":
                        result_value = result["match"]
                        
                    # Case-insensitive substring search
                    if filter_text.lower() not in result_value.lower():
                        matches_all_filters = False
                        break
                        
                if matches_all_filters:
                    filtered_results.append(result)
        
        # Sort the results if a sort column is specified
        if self.sort_column:
            try:
                key_func = None
                if self.sort_column == "File":
                    key_func = lambda r: r["file"].lower()
                elif self.sort_column == "Page":
                    key_func = lambda r: r["page"]
                elif self.sort_column == "Match":
                    key_func = lambda r: str(r["match"]).lower()
                
                filtered_results.sort(key=key_func, reverse=self.sort_reverse)
            except Exception as e:
                print(f"Error sorting: {e}")
        
        # Add sorted/filtered results to treeview
        for result in filtered_results:
            self.add_result_to_tree(result)


if __name__ == "__main__":
    root = tk.Tk()
    app = PDFRegexSearchApp(root)
    root.mainloop()
